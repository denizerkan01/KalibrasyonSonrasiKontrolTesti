import openpyxl
import os
from openpyxl.styles import Alignment
from openpyxl.styles import Font
import datetime

HEADERS = {
    "A": ("Device SN\\Test Adımları", 30),
    "B": ("Tarih - Zaman", 30),
    "C": ("Anten Testi Sonucu", 30),
    "D": ("Görünen Uydu Sayısı", 30),
    "E": ("RTK - 1 Testi Sonucu", 30),
    "F": ("RTK - 1 Testi Görünen Uydu Sayısı", 45),
    "G": ("RTK - 2 Testi Sonucu", 45),
    "H": ("RTK - 2 Testi Görünen Uydu Sayısı", 45),
    "I": ("RTK Testi Sonucu", 45),
    }

Result_Values = {

    "D": ("15'ten fazla olmalı", 30),
    "F": ("15'ten fazla olmalı", 45),
    "H": ("15'ten fazla olmalı", 45),
    }

def _ensure_headers(sheet):

    sheet.row_dimensions[1].height = 30
    for col, (text, width) in HEADERS.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}1"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

def _ensure_res_val(sheet):

    sheet.row_dimensions[2].height = 30
    for col, (text, width) in Result_Values.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}2"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

def find_row_by_sn(sheet, sn):
    row = 3
    last_match = None #yeni eklendi
    while True:
        cell_value = sheet.cell(row=row, column=1).value
        if cell_value is None:
            return last_match
        if cell_value == sn:
            #return row
            last_match = row
        row += 1

def write_result(sheet, sn, column, value, red=False):
    row = find_row_by_sn(sheet, sn)
    if row is None:
        #print(f"Could not write result — {sn} not found in column A")
        return
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    if red:
        cell.font=Font(color="FF0000")

def Anten_Sonuc(pn, sn, base_path, antenna_result, rtk_result):
    folder_path = os.path.join(base_path, "Anten Testi Sonuclari")
    os.makedirs(folder_path, exist_ok=True)
    excel_folder_path = os.path.join(folder_path, "AntenTestiSonuclari.xlsx")

    if os.path.exists(excel_folder_path):
        wb = openpyxl.load_workbook(excel_folder_path)
        sheet = wb[pn] if pn in wb.sheetnames else wb.create_sheet(pn)
        if "Sheet1" in wb.sheetnames and "Sheet1" != pn:
            default_sheet = wb["Sheet1"]
            if default_sheet.max_row == 1 and default_sheet.max_column == 1 and default_sheet["A1"].value is None:
                del wb["Sheet1"]
    else:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = pn

    _ensure_headers(sheet)  
    _ensure_res_val(sheet)

    cell = sheet.cell(row=2, column=1)
    value = "Beklenen Sonuç"
    cell.value = value
    cell.font = Font(size=14, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    
    row = 3
    while True:
        cell_value = sheet.cell(row=row, column=1).value
    
        if cell_value is None:
            sheet.cell(row=row, column=1).value = sn
            break
    
        #if cell_value == sn:
            #break
    
        row += 1
    
    current_time = datetime.datetime.now()
    
    write_result(sheet, sn, 2, current_time)

    if antenna_result["GPSSatSuccess"] == True:
        write_result(sheet, sn, 3, "Geçti") 
    else:
        write_result(sheet, sn, 3, "Kaldı", red=True)

    write_result(sheet, sn, 4, antenna_result["AntennaNumber"]) 

    if rtk_result["RTK_Port_1"] == True:
        write_result(sheet, sn, 5, "Geçti") 
    else:
        write_result(sheet, sn, 5, "Kaldı", red=True)

    write_result(sheet, sn, 6, rtk_result["Antenna_1"])

    if rtk_result["RTK_Port_2"] == True:
        write_result(sheet, sn, 7, "Geçti") 
    else:
        write_result(sheet, sn, 7, "Kaldı", red=True)

    write_result(sheet, sn, 8, rtk_result["Antenna_2"])

    if rtk_result["RTK_Success"] == True:
        write_result(sheet, sn, 9, "Geçti") 
    else:
        write_result(sheet, sn, 9, "Kaldı", red=True)

    wb.save(excel_folder_path)