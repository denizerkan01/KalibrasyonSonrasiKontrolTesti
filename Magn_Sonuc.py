import openpyxl
import os
from openpyxl.styles import Alignment
from openpyxl.styles import Font
import datetime

HEADERS = {
    "A": ("Device SN\\Test Adımları", 30),
    "B": ("Tarih - Zaman", 30),
    "C": ("Magn Testi Sonucu", 30),
    "D": ("Magn X Ortalama Değeri", 45),
    "E": ("Magn Y Ortalama Değeri", 45),
    "F": ("Magn Z Ortalama Değeri", 45),
    "G": ("Config Atılması", 30),
    "H": ("Config Alınması", 30),
    }

Result_Values = {

    "G": ("Magn Testi Başarılı ise Config Atılır", 60),
    "H": ("Magn Testi Başarılı ise Config Alınır", 60),
    }

def _ensure_headers(sheet):

    sheet.row_dimensions[1].height = 30
    for col, (text, width) in HEADERS.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}1"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

def _ensure_res_val(sheet):

    sheet.row_dimensions[2].height = 30
    for col, (text, width) in Result_Values.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}2"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

def find_row_by_sn(sheet, sn):
    row = 3
    last_match = None #yeni eklendi
    while True:
        cell_value = sheet.cell(row=row, column=1).value
        if cell_value is None:
            return last_match
        if cell_value == sn:
            #return row
            last_match = row
        row += 1

def write_result(sheet, sn, column, value, red=False):
    row = find_row_by_sn(sheet, sn)
    if row is None:
        #print(f"Could not write result — {sn} not found in column A")
        return
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    if red:
        cell.font=Font(color="FF0000")


def Magn_Sonuc(pn, sn, base_path, magn_result):
    folder_path = os.path.join(base_path, "Magn Testi Sonuclari")
    os.makedirs(folder_path, exist_ok=True)
    excel_folder_path = os.path.join(folder_path, "MagnTestiSonuclari.xlsx")

    if os.path.exists(excel_folder_path):
        wb = openpyxl.load_workbook(excel_folder_path)
        sheet = wb[pn] if pn in wb.sheetnames else wb.create_sheet(pn)
        if "Sheet1" in wb.sheetnames and "Sheet1" != pn:
            default_sheet = wb["Sheet1"]
            if default_sheet.max_row == 1 and default_sheet.max_column == 1 and default_sheet["A1"].value is None:
                del wb["Sheet1"]
    else:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = pn

    _ensure_headers(sheet)  
    _ensure_res_val(sheet)

    cell = sheet.cell(row=2, column=1)
    value = "Beklenen Sonuç"
    cell.value = value
    cell.font = Font(size=14, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    
    row = 3
    while True:
        cell_value = sheet.cell(row=row, column=1).value
    
        if cell_value is None:
            sheet.cell(row=row, column=1).value = sn
            break
    
        #if cell_value == sn:
            #break
    
        row += 1
    
    current_time = datetime.datetime.now()
    
    write_result(sheet, sn, 2, current_time)

    if magn_result["MagnSuccess"] == True:
        write_result(sheet, sn, 3, "Geçti") 
    else:
        write_result(sheet, sn, 3, "Kaldı", red=True)

    write_result(sheet, sn, 4, magn_result["Mean_MagnX"]) 
    write_result(sheet, sn, 5, magn_result["Mean_MagnY"]) 
    write_result(sheet, sn, 6, magn_result["Mean_MagnZ"]) 

    if magn_result["MagnSuccess"] == True:
        write_result(sheet, sn, 7, "Atıldı") 
    else:
        write_result(sheet, sn, 7, "Atılmadı", red=True)

    if magn_result["MagnSuccess"] == True:
        write_result(sheet, sn, 8, "Alındı") 
    else:
        write_result(sheet, sn, 8, "Alınmadı", red=True)

    wb.save(excel_folder_path)